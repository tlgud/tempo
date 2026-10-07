#!/usr/bin/env python3
"""HB-analysis-06-6 exact RPM, SSI, lot-stat, and template tests.

The analyzer is the file named by HB_ANALYZER_UNDER_TEST.
"""
import hashlib
import importlib.util
import os
import shutil
import subprocess
import sys
from pathlib import Path

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent
EXPECTED_VERSION = "HB-analysis-06-6"
EXPECTED_QC_LAYER_VERSION = "LotQC-2.8"
CODE = "A_SSI_QC_0P5_5"


def _target_from_env():
    raw_target = os.environ.get("HB_ANALYZER_UNDER_TEST")
    if not raw_target:
        raise RuntimeError("HB_ANALYZER_UNDER_TEST must contain the exact script path")
    target = Path(raw_target).expanduser().resolve()
    if not target.is_file():
        raise FileNotFoundError(f"Analyzer under test does not exist: {target}")
    return target


TARGET = _target_from_env()
TARGET_SHA256 = hashlib.sha256(TARGET.read_bytes()).hexdigest()
_LOAD_SEQ = 0
PROTOCOL = [0.0, 0.2, 0.35, 0.5, 0.7, 1.0, 1.5, 2.0, 3.0, 4.0, 5.0]
DOWN = [4.0, 3.0, 2.0, 1.5, 1.0, 0.7, 0.5, 0.35, 0.2]


def load_mod(path: Path):
    global _LOAD_SEQ
    path = Path(path).expanduser().resolve()
    _LOAD_SEQ += 1
    modname = f"hb06_under_test_{_LOAD_SEQ}"
    spec = importlib.util.spec_from_file_location(modname, path)
    mod = importlib.util.module_from_spec(spec)
    sys.modules[modname] = mod
    spec.loader.exec_module(mod)
    mod._verified_path = str(path)
    mod._verified_sha256 = hashlib.sha256(path.read_bytes()).hexdigest()
    return mod


M = load_mod(TARGET)


def test_00_target_identity():
    print(M._verified_path)
    print(M.VERSION)
    print(M.QC_LAYER_VERSION)
    print(M._verified_sha256)
    assert TARGET.is_file()
    assert Path(M._verified_path).resolve() == TARGET.resolve()
    assert M.VERSION == EXPECTED_VERSION
    assert M.QC_LAYER_VERSION == EXPECTED_QC_LAYER_VERSION
    assert len(M._verified_sha256) == 64
    assert M._verified_sha256 == TARGET_SHA256
    expected_sha = os.environ.get("HB_ANALYZER_EXPECTED_SHA256")
    if expected_sha:
        assert M._verified_sha256 == expected_sha


def test_00_specified_path_is_verified_without_basename(tmp_path):
    original_name = TARGET.name
    payload = TARGET.read_bytes()
    etas = {2: 240, 4: 180, 8: 150, 10: 120, 11: 100}
    for name in ("analyzer.py", "사용자_지정_이름.py", "rheology-final.py"):
        copied = tmp_path / name
        copied.write_bytes(payload)
        mod = load_mod(copied)
        assert Path(mod._verified_path).resolve() == copied.resolve()
        assert mod.VERSION == EXPECTED_VERSION
        assert mod.QC_LAYER_VERSION == EXPECTED_QC_LAYER_VERSION
        assert mod._verified_sha256 == TARGET_SHA256
        csv_path = write_csv(tmp_path / f"data-{copied.stem}.csv", steps_of(PROTOCOL), etas=etas)
        _reps, qc_s, qc_i, _rs, _ri = analyze_with(mod, csv_path)
        assert qc_s["values"]["A_SSI_FIELD_0P2_2"] == pytest.approx(1.6)
        assert qc_s["values"]["A_SSI_QC_0P5_5"] == pytest.approx(1.8)
    assert TARGET.name == original_name
    assert TARGET.read_bytes() == payload


def write_csv(path, steps, hold=60, holds=None, start=1, etas=None):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    holds = holds or {}
    etas = etas or {}
    lines = [
        "Test Start,2026-09-20 10:00:00,Instrument,DVNext,Spindle,CP-52",
        "Temperature Control,No,Sample,synthetic",
        "",
        "Step,Speed,Temperature,Data Collection",
        "hdr2",
    ]
    for offset, (rpm, _tq) in enumerate(steps):
        step = start + offset
        h = int(holds.get(step, hold if rpm else 600))
        lines.append(f"{step},{rpm},25,Time = 00:{h // 60:02d}:{h % 60:02d}")
    lines += ["", "Step,Point,Time,Viscosity,Torque,Speed,Shear Stress,Shear Rate,Temperature,Density,Accuracy", "units"]
    t = 0
    pt = 0
    for offset, (rpm, tq) in enumerate(steps):
        step = start + offset
        h = int(holds.get(step, hold if rpm else 600))
        g = 2.0 * float(rpm)
        torque = float(tq)
        eta = etas.get(step)
        for _k in range(1, h + 1):
            t += 1
            pt += 1
            visc = (torque * 10.0 / g * 1000.0) if (eta is None and g) else (float(eta) if eta is not None else 0.0)
            tau = torque * 10.0
            lines.append(
                f"{step},{pt},{t},{visc:.4f},{torque:.4f},{rpm},{tau * 10:.4f},{g:.3f},25.00,1.05,{max(visc * 0.01, 0.1):.4f}"
            )
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return path


def analyze_with(mod, path, policy=None):
    policy = policy or mod.POLICY_INCLUSIVE
    mod.set_active_policy(policy)
    df, holds, meta = mod.parse_dvnext_csv(path)
    reps = mod.step_representatives(df, holds, policy=policy)
    strict_view = mod.channel_view(reps, "strict")
    incl_view = mod.channel_view(reps, "inclusive")
    result_s = mod.analyze_step_flow(strict_view)
    result_i = mod.analyze_step_flow(incl_view)
    qc_s = mod.compute_a_qc(df, holds, meta, strict_view, result_s, {}, [], "VALID", 10, 95, channel="strict", policy=policy)
    qc_i = mod.compute_a_qc(df, holds, meta, incl_view, result_i, {}, [], "VALID", 10, 95, channel="inclusive", policy=policy)
    return reps, qc_s, qc_i, result_s, result_i


def analyze(path, policy=None):
    return analyze_with(M, path, policy)


def steps_of(rpms, torque=20.0):
    return [(rpm, 0.0 if rpm == 0 else torque) for rpm in rpms]


def finite(x):
    try:
        return x is not None and np.isfinite(float(x))
    except (TypeError, ValueError):
        return False


def pair_ids(pairs):
    return sorted(p["pair_id"] for p in pairs)


def _rpm_frame(rows):
    return pd.DataFrame(rows)


def _up_row(step, rpm, eta, included=True, n_valid=5, value_class=None, turnaround=False, direction="Up"):
    klass = value_class if value_class is not None else (M.VC_STRICT if included else M.VC_NO)
    return {
        "Step": step, "RPM": rpm, "gamma": 2.0 * rpm, "eta_cP": eta, "tau_Pa": eta * rpm / 1000.0,
        "n_valid": n_valid if included else 0, "moving_flag": True, "step_type": "MOVING",
        "sequence_direction": "Turnaround" if turnaround else direction,
        "dir": "Turnaround" if turnaround else direction,
        "used_in_up_branch": direction == "Up" or turnaround, "used_in_down_branch": direction == "Down" or turnaround,
        "turnaround_shared": turnaround, "included_in_strict": included,
        "value_class_strict": klass, "data_status": "OK" if included else "EXCLUDED",
        "acceptance_status": "ELIGIBLE" if included else "EXCLUDED",
    }


def test_01_branch_leak_blocked(tmp_path):
    path = write_csv(tmp_path / "HT-B1-01.csv", steps_of([0.0, 0.2, 2.0, 5.0, 0.5]))
    reps, qc_s, _qi, _rs, _ri = analyze(path)
    st = qc_s["status"][CODE]
    down = reps[(reps["RPM"] - 0.5).abs() < 1e-6]
    assert len(down) == 1
    assert down.iloc[0]["sequence_direction"] == "Down"
    assert not finite(st["value"])
    assert st["value_class"] == M.VC_NO
    assert st["numerator_step"] is None
    assert int(down.iloc[0]["Step"]) in st["numerator_candidate_steps"]
    assert "Up 후보 없음" in st["reason"]


def test_02_turnaround_not_used_as_low_point(tmp_path):
    path = write_csv(tmp_path / "HT-B2-01.csv", steps_of([0.0, 0.2, 0.5, 0.2]))
    reps, _qs, _qi, _rs, _ri = analyze(path)
    picked = M.select_exact_rpm_step(reps, 0.5, "strict", branch="up", allow_turnaround=False)
    assert picked["found"] is False
    assert picked["selected_step"] is None
    assert "shared turnaround" in picked["exclusion_reason"]
    assert "allow_turnaround=False" in picked["exclusion_reason"]


def test_03_qc_denominator_uses_shared_turnaround(tmp_path):
    path = write_csv(tmp_path / "HT-B3-01.csv", steps_of(PROTOCOL + DOWN), etas={4: 180, 11: 100})
    reps, _qs, qc_i, _rs, _ri = analyze(path)
    five = reps[(reps["RPM"] - 5).abs() < 1e-6]
    assert len(five) == 1
    assert bool(five.iloc[0]["turnaround_shared"]) is True
    st = qc_i["status"][CODE]
    assert finite(st["value"])
    assert st["denominator_step"] == int(five.iloc[0]["Step"])
    assert st["turnaround_shared"] is True


def test_04_unique_eligible_among_measured_duplicates():
    frame = _rpm_frame([
        _up_row(3, 0.5, 180, included=False),
        _up_row(4, 0.5, 180, included=True),
        _up_row(8, 5.0, 100, included=True),
    ])
    picked = M.select_exact_rpm_step(frame, 0.5, "strict", branch="up", allow_turnaround=False)
    assert picked["found"] is True
    assert picked["selected_step"] == 4
    assert picked["data_status"] == "REVIEW"
    assert picked["duplicate_count"] == 2
    assert "3" in picked["exclusion_reason"]


def test_05_two_eligible_candidates_are_not_picked():
    frame = _rpm_frame([
        _up_row(3, 0.5, 180, included=True),
        _up_row(4, 0.5, 170, included=True),
        _up_row(8, 5.0, 100, included=True),
    ])
    picked = M.select_exact_rpm_step(frame, 0.5, "strict", branch="up", allow_turnaround=False)
    metric = M.compute_exact_ssi_metric(frame, "strict", CODE, 0.5, 5.0, allow_high_turnaround=True)
    assert picked["found"] is False
    assert picked["ambiguous"] is True
    assert picked["eligible_candidate_steps"] == [3, 4]
    assert not finite(metric["value"])
    assert metric["numerator_step"] is None
    assert metric["numerator_candidate_steps"] == [3, 4]


def test_06_official_ssi_numbers(tmp_path):
    etas = {2: 240, 4: 180, 8: 150, 10: 120, 11: 100}
    path = write_csv(tmp_path / "HT-N6-01.csv", steps_of(PROTOCOL), etas=etas)
    _reps, qc_s, qc_i, _rs, _ri = analyze(path)
    assert qc_s["values"]["A_SSI_FIELD_0P2_2"] == pytest.approx(1.6)
    assert qc_i["values"][CODE] == pytest.approx(1.8)
    field = qc_i["status"]["A_SSI_FIELD_0P2_2"]
    qc = qc_i["status"][CODE]
    assert field["numerator_rpm"] == pytest.approx(0.2)
    assert field["denominator_rpm"] == pytest.approx(2.0)
    assert qc["numerator_rpm"] == pytest.approx(0.5)
    assert qc["denominator_rpm"] == pytest.approx(5.0)
    assert field["selection_rule"] == M.FIELD_SSI_SELECTION_RULE
    assert qc["selection_rule"] == M.QC_SSI_SELECTION_RULE


def test_07_ineligible_5rpm_is_not_replaced_by_4rpm(tmp_path):
    steps = []
    for rpm in PROTOCOL:
        steps.append((rpm, 100.0 if rpm == 5 else 20.0))
    path = write_csv(tmp_path / "HT-N7-01.csv", steps, etas={4: 180, 10: 120, 11: 100})
    _reps, qc_s, qc_i, _rs, _ri = analyze(path)
    strict = qc_s["status"][CODE]
    incl = qc_i["status"][CODE]
    assert not finite(strict["value"])
    assert strict["denominator_step"] is None
    assert any(True for _s in strict["denominator_candidate_steps"])
    assert "4.0 rpm으로 대체하지 않음" in strict["reason"]
    assert finite(incl["value"])
    assert incl["denominator_rpm"] == pytest.approx(5.0)
    assert incl["value_class"] in (M.VC_HIGH, M.VC_BOTH, M.VC_LOW)


def test_08_legacy_migration_uses_provenance():
    mapped_qc = M.canonicalize_metric_record({
        "metric_code": "A_SSI", "numerator_rpm": 0.5, "denominator_rpm": 5.0,
        "direction": "Up", "source_type": "measured",
    })
    mapped_field = M.canonicalize_metric_record({
        "metric_code": "A_SSI", "numerator_rpm": 0.2, "denominator_rpm": 2.0,
        "direction": "Up", "source_type": "measured",
    })
    unknown = M.canonicalize_metric_record({"metric_code": "A_SSI"})
    legacy_4 = M.canonicalize_metric_record({
        "metric_code": "A_SSI", "numerator_rpm": 0.5, "denominator_rpm": 4.0,
        "direction": "Up", "source_type": "measured",
    })
    assert mapped_qc["canonical_metric_code"] == CODE
    assert mapped_field["canonical_metric_code"] == "A_SSI_FIELD_0P2_2"
    assert unknown["canonical_metric_code"] == "A_SSI_LEGACY_UNKNOWN"
    assert unknown["included_in_lot_stats"] is False
    assert unknown["acceptance_status"] == "EXCLUDED"
    assert legacy_4["canonical_metric_code"] == "A_SSI_LEGACY_UNKNOWN"
    with pytest.raises(ValueError, match="code alone"):
        M.metric_alias("A_SSI")


def _both_rec(value=8.0):
    st = {"status": "CENSORED", "data_status": "CENSORED", "value_class": M.VC_BOTH,
          "source_type": "measured", "acceptance_status": "EXCLUDED", "value": value}
    return {"repeat": "01", "values": {CODE: value}, "status": {CODE: st},
            "values_incl": {CODE: value}, "status_incl": {CODE: st}}


def test_09_vc_both_excluded_from_mean():
    policy = M.replace(M.POLICY_INCLUSIVE, include_censored_in_lot_stats=False)
    decision = M.lot_inclusion_decision(
        metric_code=CODE, value=8.0, channel="inclusive", data_status="CENSORED",
        value_class=M.VC_BOTH, acceptance_status="EXCLUDED", source_type="measured",
        pair_id=None, policy=policy,
    )
    stats = M.lot_metric_stats([_both_rec()], CODE, {"method": "A"}, "inclusive", policy)
    assert decision["included"] is False
    assert not finite(decision["stat_value"])
    assert stats["n"] == 0
    assert stats["n_censored"] == 1
    assert stats["n_low_ref"] == 1
    assert stats["n_high_ref"] == 1


def test_10_vc_both_included_keeps_acceptance_excluded():
    policy = M.replace(M.POLICY_INCLUSIVE, include_censored_in_lot_stats=True)
    decision = M.lot_inclusion_decision(
        metric_code=CODE, value=8.0, channel="inclusive", data_status="CENSORED",
        value_class=M.VC_BOTH, acceptance_status="EXCLUDED", source_type="measured",
        pair_id=None, policy=policy,
    )
    stats = M.lot_metric_stats([_both_rec()], CODE, {"method": "A"}, "inclusive", policy)
    assert decision["included"] is True
    assert decision["stat_value"] == pytest.approx(8.0)
    assert stats["acceptance_status"] == "EXCLUDED"
    assert stats["n"] == 1


def _mix_rec(repeat, value, status, klass):
    st = {"status": status, "data_status": status, "value_class": klass, "source_type": "measured",
          "reason": "",
          "acceptance_status": "EXCLUDED" if klass in M.TORQUE_REFERENCE_CLASSES else "ELIGIBLE",
          "value": value}
    return {"lot": "MIX", "repeat": repeat, "file": f"{repeat}.csv", "test_start": "", "route": "INCLUSIVE",
            "data_qc": "VALID", "method": "A", "audit_rows": [],
            "values": {CODE: value}, "status": {CODE: st},
            "values_incl": {CODE: value}, "status_incl": {CODE: st}}


def _mix_records():
    return [
        _mix_rec("01", 1.0, "OK", M.VC_STRICT),
        _mix_rec("02", 3.0, "OK", M.VC_STRICT),
        _mix_rec("03", 2.0, "EXCEPTION", M.VC_EXC),
        _mix_rec("04", 4.0, "CENSORED", M.VC_LOW),
        _mix_rec("05", 5.0, "CENSORED", M.VC_HIGH),
        _mix_rec("06", 6.0, "CENSORED", M.VC_BOTH),
        _mix_rec("07", np.nan, "NO VALUE", M.VC_NO),
    ]


def _soffice():
    found = shutil.which("soffice")
    if found:
        return Path(found)
    for candidate in (Path(r"C:\Program Files\LibreOffice\program\soffice.exe"),
                      Path(r"C:\Program Files (x86)\LibreOffice\program\soffice.exe")):
        if candidate.exists():
            return candidate
    pytest.fail("LibreOffice soffice.exe 가 없어 재계산 검사를 실행할 수 없습니다.")


def _recalc(tmp_path, workbook):
    profile = tmp_path / f"lo-{workbook.stem}"
    converted = tmp_path / f"conv-{workbook.stem}"
    converted.mkdir()
    subprocess.run([
        str(_soffice()), "--headless", "--norestore", "--nolockcheck", "--nologo",
        f"-env:UserInstallation={profile.resolve().as_uri()}",
        "--convert-to", "xlsx", "--outdir", str(converted), str(workbook),
    ], check=True, timeout=180)
    return load_workbook(next(converted.glob("*.xlsx")), data_only=True)


def _excel_stats(wb, channel):
    ws = wb["Metric_Statistics"]
    header = {ws.cell(5, col).value: col for col in range(2, 30)}
    found = None
    for row in range(6, ws.max_row + 1):
        if ws.cell(row, header["Code"]).value == CODE:
            found = row
            break
    assert found is not None
    if channel == "strict":
        keys = {"n": "n Strict", "mean": "Strict Mean", "sd": "Strict SD", "min": "Strict Min",
                "max": "Strict Max", "range": "Strict Range"}
    else:
        keys = {"n": "n Inclusive", "mean": "Inclusive Mean", "sd": "Inclusive SD",
                "min": "Inclusive Min", "max": "Inclusive Max", "range": "Inclusive Range"}
    return {name: ws.cell(found, header[label]).value for name, label in keys.items()}


def _same(python_value, excel_value):
    if not finite(python_value):
        return excel_value in (None, "")
    return abs(float(python_value) - float(excel_value)) <= 1e-9


def test_11_python_excel_stats_match(tmp_path):
    recs = _mix_records()
    for flag in (False, True):
        policy = M.replace(
            M.POLICY_INCLUSIVE, include_censored_in_lot_stats=flag,
            name="INCL_ON" if flag else "INCL_OFF",
        )
        M.set_active_policy(policy)
        folder = tmp_path / policy.name
        folder.mkdir()
        out, _stats = M.write_lot_qc_workbook("MIX", recs, [], [], folder, policy=policy)
        wb = _recalc(tmp_path, Path(out))
        for channel in ("strict", "inclusive"):
            py = M.lot_metric_stats(recs, CODE, {"method": "A"}, channel, policy)
            ex = _excel_stats(wb, channel)
            assert int(ex["n"] or 0) == py["n"]
            assert _same(py["mean"], ex["mean"])
            assert _same(py["sd"], ex["sd"])
            assert _same(py["min"], ex["min"])
            assert _same(py["max"], ex["max"])
            assert _same(py["range"], ex["range"])
        incl = M.lot_metric_stats(recs, CODE, {"method": "A"}, "inclusive", policy)
        assert (incl["n"] == 6) if flag else (incl["n"] == 3)
        for rec in recs:
            row = rec["status_incl"][CODE]
            if row["value_class"] in M.TORQUE_REFERENCE_CLASSES:
                assert row["acceptance_status"] == "EXCLUDED"


def test_12_dynamic_decade_up_only(tmp_path):
    path = write_csv(tmp_path / "HT-D12-01.csv", steps_of(PROTOCOL))
    _reps, _qs, qc_i, _rs, _ri = analyze(path)
    ids = pair_ids(qc_i["decade_pairs"])
    assert ids == ["SSI_DECADE_UP_0P2_2", "SSI_DECADE_UP_0P5_5"]
    links = {p["pair_id"]: p["official_metric_link"] for p in qc_i["decade_pairs"]}
    assert links["SSI_DECADE_UP_0P2_2"] == "A_SSI_FIELD_0P2_2"
    assert links["SSI_DECADE_UP_0P5_5"] == CODE
    assert all(p["direction"] == "up" for p in qc_i["decade_pairs"])


def test_13_changed_protocol_keeps_official_rpm(tmp_path):
    path = write_csv(tmp_path / "HT-D13-01.csv", steps_of([0.0, 0.3, 0.5, 1.0, 2.0, 3.0, 5.0]))
    _reps, qc_s, qc_i, _rs, _ri = analyze(path)
    ids = pair_ids(qc_i["decade_pairs"])
    assert ids == ["SSI_DECADE_UP_0P3_3", "SSI_DECADE_UP_0P5_5"]
    assert not finite(qc_s["values"]["A_SSI_FIELD_0P2_2"])
    assert finite(qc_i["values"][CODE])
    assert qc_i["status"][CODE]["denominator_rpm"] == pytest.approx(5.0)
    assert all(p["official_metric_link"] != "A_SSI_FIELD_0P2_2" for p in qc_i["decade_pairs"])


def test_14_down_decade_is_optional_and_separate(tmp_path):
    path = write_csv(tmp_path / "HT-D14-01.csv", steps_of(PROTOCOL + DOWN))
    off = M.replace(M.POLICY_INCLUSIVE, ssi_include_down=False)
    on = M.replace(M.POLICY_INCLUSIVE, ssi_include_down=True)
    _r, _s, qc_off, _a, _b = analyze(path, off)
    _r, _s, qc_on, _a, _b = analyze(path, on)
    assert {p["direction"] for p in qc_off["decade_pairs"]} == {"up"}
    assert {p["direction"] for p in qc_on["decade_pairs"]} == {"up", "down"}
    up_ids = {p["pair_id"] for p in qc_on["decade_pairs"] if p["direction"] == "up"}
    dn_ids = {p["pair_id"] for p in qc_on["decade_pairs"] if p["direction"] == "down"}
    assert up_ids.isdisjoint(dn_ids)
    assert all(p["official_metric_link"] == "" for p in qc_on["decade_pairs"] if p["direction"] == "down")
    grouped = {}
    for pair in qc_on["decade_pairs"]:
        grouped.setdefault(pair["pair_id"], set()).add(pair["direction"])
    assert all(len(v) == 1 for v in grouped.values())


def test_15_unsupported_cli_options_raise():
    parser = M.parser()
    with pytest.raises(ValueError, match="pair identity must be preserved"):
        M.policy_from_args(parser.parse_args(["--ssi-pool-different-pairs"]))
    with pytest.raises(NotImplementedError, match="Estimated decade SSI"):
        M.policy_from_args(parser.parse_args(["--ssi-allow-interpolation"]))
    with pytest.raises(NotImplementedError, match="Estimated decade SSI"):
        M.policy_from_args(parser.parse_args(["--ssi-allow-model-estimation"]))


def test_16_high_shear_stays_apart_from_official_5rpm(tmp_path):
    path = write_csv(tmp_path / "HT-H16-01.csv", steps_of([0.0, 0.2, 0.5, 2.0, 5.0, 7.0]))
    _reps, qc_s, _qi, _rs, _ri = analyze(path)
    assert qc_s["status"]["A_TAU_H_UP"]["selected_high_rpm"] == pytest.approx(7.0)
    assert qc_s["status"]["A_ETA_HIGH"]["selected_high_rpm"] == pytest.approx(7.0)
    assert qc_s["status"][CODE]["denominator_rpm"] == pytest.approx(5.0)
    assert qc_s["status"]["A_SSI_FIELD_0P2_2"]["denominator_rpm"] == pytest.approx(2.0)


def test_17_up_only_has_no_down_or_hysteresis(tmp_path):
    path = write_csv(tmp_path / "HT-U17-01.csv", steps_of(PROTOCOL))
    reps, _qs, _qi, result, _ri = analyze(path)
    rest = reps[reps["step_type"] == "REST_STABILIZATION"]
    moving = reps[reps["moving_flag"]]
    end = moving.sort_values("RPM").iloc[-1]
    assert len(rest) == 1
    assert set(moving["sequence_direction"]) == {"Up"}
    assert abs(float(end["RPM"]) - 5.0) < 1e-6
    assert bool(end["turnaround_shared"]) is False
    assert str(result["fits"]["down_full"]["status"]).startswith("INVALID")
    assert str(result["fits"]["up_common"]["status"]).startswith("INVALID")
    assert not finite(result["relative_area_pct"])


def test_18_real_down_branch_shares_one_turnaround(tmp_path):
    path = write_csv(tmp_path / "HT-U18-01.csv", steps_of(PROTOCOL + DOWN))
    reps, _qs, _qi, result, result_i = analyze(path)
    five = reps[(reps["RPM"] - 5).abs() < 1e-6]
    assert len(five) == 1
    assert bool(five.iloc[0]["turnaround_shared"]) is True
    for fit_key in ("up_full", "down_full", "up_common", "down_common"):
        assert result["fits"][fit_key]["n_points"] > 0
    assert finite(result["relative_area_pct"])
    assert finite(result_i["relative_area_pct"])


def test_19_template_schema(tmp_path):
    good = load_workbook(ROOT / "SST_Gel_StepFlow_Template.xlsx")
    M.validate_template_schema(good, "A", ROOT / "SST_Gel_StepFlow_Template.xlsx")
    good.close()
    path = write_csv(tmp_path / "HT-T19-01.csv", steps_of(PROTOCOL))
    M.set_active_policy(M.POLICY_INCLUSIVE)
    out = tmp_path / "stepflow.xlsx"
    M.process_enhanced(path, "A", ROOT / "SST_Gel_StepFlow_Template.xlsx", out, 10, 0.25, 10, 95)
    assert out.exists()
    broken = tmp_path / "broken-header.xlsx"
    shutil.copyfile(ROOT / "SST_Gel_StepFlow_Template.xlsx", broken)
    wb = load_workbook(broken)
    wb["Step_Setup"]["B5"] = "WRONG"
    wb.save(broken)
    wb.close()
    reopened = load_workbook(broken)
    with pytest.raises(M.TemplateSchemaError, match="expected='Step'"):
        M.validate_template_schema(reopened, "A", broken)
    missing = tmp_path / "missing-sheet.xlsx"
    shutil.copyfile(ROOT / "SST_Gel_3ITT_Template.xlsx", missing)
    wb = load_workbook(missing)
    del wb["Recovery"]
    wb.save(missing)
    wb.close()
    reopened = load_workbook(missing)
    with pytest.raises(M.TemplateSchemaError, match="Recovery"):
        M.validate_template_schema(reopened, "B", missing)


def test_20_summary_lists_each_official_ssi_once(tmp_path):
    path = write_csv(tmp_path / "HT-S20-01.csv", steps_of(PROTOCOL), etas={2: 240, 4: 180, 8: 150, 11: 100})
    _reps, qc_s, qc_i, _rs, _ri = analyze(path)
    rec = dict(qc_s)
    rec.update({"lot": "S20", "repeat": "01", "file": path.name, "method": "A", "test_start": "",
                "route": "INCLUSIVE", "data_qc": "VALID", "audit_rows": [],
                "values_incl": qc_i["values"], "status_incl": qc_i["status"],
                "decade_pairs_strict": qc_s["decade_pairs"], "decade_pairs_incl": qc_i["decade_pairs"]})
    out, _stats = M.write_lot_qc_workbook("S20", [rec], [], [], tmp_path, policy=M.POLICY_INCLUSIVE)
    wb = load_workbook(out)
    codes = [wb["Lot_QC_Summary"].cell(r, 3).value for r in range(1, wb["Lot_QC_Summary"].max_row + 1)]
    assert codes.count("A_SSI_FIELD_0P2_2") == 1
    assert codes.count(CODE) == 1
    for banned in ("A_SSI", "A_SSI_DYNAMIC", "A_SSI_REF_0P5_4", "A_SSI_LEGACY_UNKNOWN", "A_SSI_DECADE_ESTIMATED"):
        assert banned not in codes
    metric_codes = [wb["Metric_Statistics"].cell(r, 2).value for r in range(6, wb["Metric_Statistics"].max_row + 1)]
    assert metric_codes.count("A_SSI_FIELD_0P2_2") == 1
    assert metric_codes.count(CODE) == 1
    assert not any(str(c).startswith("SSI_DECADE_") for c in metric_codes)
    decade_ids = [wb["DecadeSSI"].cell(r, 6).value for r in range(6, wb["DecadeSSI"].max_row + 1)]
    assert "SSI_DECADE_UP_0P2_2" in decade_ids
    assert "SSI_DECADE_UP_0P5_5" in decade_ids
